"""Extract plain text from uploaded resumes / job descriptions."""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class UnsupportedFileError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # a single malformed page should not kill the upload
            pages.append("")
    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Return cleaned text for a supported document, raising on unsupported types."""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"Unsupported file type '{ext or filename}'. Upload a PDF, DOCX, TXT or MD file, or paste the text."
        )
    if ext == ".pdf":
        raw = _pdf_text(data)
    elif ext == ".docx":
        try:
            raw = _docx_text(data)
        except Exception as exc:  # python-docx raises a variety of zip/xml errors
            raise UnsupportedFileError(f"Could not read '{filename}' as a Word document: {exc}") from exc
    else:
        raw = _plain_text(data)
    return _clean(raw)


def is_pdf(filename: str) -> bool:
    return Path(filename or "").suffix.lower() == ".pdf"
