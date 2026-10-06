import io

import docx
import pytest

from app.agent.pipeline import plan_jobs
from app.agent.plan import DEPTHS, PrepContext, chunk_count, split_by_level
from app.knowledge_base import TOPICS, select_topics, topic_by_name
from app.knowledge_base.matching import extract_skills, skill_gap
from app.models import QUESTION_SECTIONS, dedupe, normalize_items, normalize_topics
from app.parsing import UnsupportedFileError, extract_text


def _minimal_pdf(text: str) -> bytes:
    """A tiny valid single-page PDF containing `text` (no external tools needed)."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


# ----------------------------------------------------------------------- parsing
def test_extract_text_from_txt_docx_and_pdf():
    assert extract_text("cv.txt", "Python  developer\r\n\r\n\r\nSQL".encode()) == "Python developer\n\nSQL"

    document = docx.Document()
    document.add_paragraph("Senior Data Scientist")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "SQL"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("resume.docx", buffer.getvalue())
    assert "Senior Data Scientist" in text and "Python | SQL" in text

    assert "Hello Interview" in extract_text("resume.pdf", _minimal_pdf("Hello Interview"))


def test_extract_text_rejects_unknown_types():
    with pytest.raises(UnsupportedFileError):
        extract_text("photo.png", b"\x89PNG")


# ------------------------------------------------------------------------ models
def test_normalize_items_canonicalises_and_drops_broken_items():
    items = normalize_items(
        "theory",
        [
            {"question": "Explain overfitting in machine learning?", "level": "Hard", "hot": "yes", "key_points": "a\nb"},
            {"question": "short"},  # too short -> dropped
            "not a dict",
        ],
    )
    assert len(items) == 1
    assert items[0]["level"] == "advanced" and items[0]["hot"] is True
    assert items[0]["key_points"] == ["a", "b"] and items[0]["id"].startswith("th-")

    quiz = normalize_items(
        "quiz",
        [
            {"question": "Which is immutable in Python?", "options": ["list", "tuple", "dict", "set"], "correct_index": 1},
            {"question": "Bad index question here?", "options": ["a", "b"], "correct_index": 5},
            {"question": "Duplicate options question?", "options": ["a", "a", "b"], "correct_index": 0},
        ],
    )
    assert len(quiz) == 1 and quiz[0]["option_explanations"] == ["", "", "", ""]


def test_dedupe_and_topics():
    existing = [{"question": "What is the difference between a list and a tuple in Python?"}]
    new = [
        {"question": "What's the difference between a tuple and a list in Python?"},
        {"question": "Explain Python generators and yield."},
    ]
    assert [q["question"] for q in dedupe(existing, new)] == ["Explain Python generators and yield."]
    topics = normalize_topics([{"name": "SQL", "weight": 0.8, "priority": "critical"}, "SQL", {"name": "Python"}])
    assert [t["name"] for t in topics] == ["SQL", "Python"]
    assert topics[0]["weight"] == 80 and topics[0]["priority"] == "high"


# -------------------------------------------------------------------------- plan
@pytest.mark.parametrize("years", [0, 1.5, 4, 7, 15])
@pytest.mark.parametrize("total", [1, 2, 3, 10, 24, 36])
def test_split_by_level_sums_to_total(years, total):
    split = split_by_level(total, years)
    assert sum(split.values()) == total
    if total >= 3:
        assert all(v >= 1 for v in split.values())


def test_seniority_shifts_the_level_mix():
    junior, senior = split_by_level(30, 0.5), split_by_level(30, 12)
    assert junior["beginner"] > senior["beginner"] and senior["advanced"] > junior["advanced"]


def test_chunking_and_job_plan_cover_every_target():
    assert chunk_count(11, 12) == [11] and chunk_count(25, 12) == [9, 8, 8] and chunk_count(0, 5) == []
    for depth, targets in DEPTHS.items():
        ctx = PrepContext("s", "Acme", "Analyst", 4, depth, "resume", "jd")
        topics = [f"T{i}" for i in range(targets["topics"])]
        jobs = plan_jobs(ctx, topics)
        for section in QUESTION_SECTIONS:
            assert sum(j.count for j in jobs if j.section == section) == targets[section]
        assert sorted(t for j in jobs if j.section == "revision" for t in j.topics) == sorted(topics)
        # every depth yields at least 50 questions
        assert sum(targets[s] for s in QUESTION_SECTIONS) >= 50


# ------------------------------------------------------------- knowledge base
def test_every_knowledge_base_item_is_valid():
    for topic in TOPICS:
        for section in QUESTION_SECTIONS:
            raw = topic[section]
            assert len(normalize_items(section, raw)) == len(raw), (topic["key"], section)
            assert {i["level"] for i in raw} == {"beginner", "intermediate", "advanced"}, (topic["key"], section)
        note = normalize_items("revision", [{"topic": topic["name"], **topic["revision"]}])
        assert note and note[0]["key_concepts"] and note[0]["explanation_md"]


def test_topic_selection_and_skill_gap():
    jd = "Senior Data Scientist: Python, SQL, machine learning, PySpark on Databricks, LLMs and RAG, Azure or AWS."
    resume = "Built XGBoost churn models in Python; SQL on Snowflake; RAG assistant with Azure OpenAI."
    names = [e["topic"]["name"] for e in select_topics(jd, resume, "Senior Data Scientist", "", 8)]
    assert len(names) == 8 and "Behavioural & HR" in names
    assert {"Python", "SQL & Databases", "Machine Learning", "NLP & Generative AI"} <= set(names)
    matched, missing = skill_gap(extract_skills(jd), extract_skills(resume))
    assert "Machine Learning" in matched  # implied by XGBoost
    assert "AWS" not in missing  # Azure covers the "Azure or AWS" requirement
    assert "Spark" in missing
    assert topic_by_name("sql")["key"] == "sql" and topic_by_name("Kafka streaming")["key"] == "data_engineering"
